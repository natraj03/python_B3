import  openpyxl

file_obj = openpyxl.load_workbook("practicefile_2.xlsx")


def test_file_mult():
    my_sheet = file_obj.active
    print("my_sheet",my_sheet)
    data_sheet = []
    for row in my_sheet.iter_rows(values_only=True):
        # print("row",row)
        # data_sheet.append(row)
        var1 = row[0]
        var2 = row[1]
        var3 = row[2]
        var4 = row[3]
        print(f"mult result {var1} * {var2} = {var1 * var2}")
        if var3 != None:
            print(f"text add string result {var3} + {var4} = {var3 + var4}")

    # print("data_sheet",data_sheet)


def test_file_add():
    my_sheet = file_obj.active
    print("my_sheet",my_sheet)
    data_sheet = []
    for row in my_sheet.iter_rows(values_only=True):
        # print("row",row)
        # data_sheet.append(row)
        var1 = row[0]
        var2 = row[1]
        print(f"add result {var1} + {var2} = {var1 + var2}")

def test_file_sub():
    my_sheet = file_obj.active
    print("my_sheet",my_sheet)
    data_sheet = []
    for row in my_sheet.iter_rows(values_only=True):
        # print("row",row)
        # data_sheet.append(row)
        var1 = row[0]
        var2 = row[1]
        print(f"sub result {var1} - {var2} = {var1 - var2}")





